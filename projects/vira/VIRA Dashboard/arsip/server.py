"""
VIRA Performance & CRM Dashboard — Server (server.py)
Handles API endpoints, database.json persistence, and Excel sheet updates.
"""

import http.server
import socketserver
import json
import os
import sys
import openpyxl
import mimetypes
from datetime import datetime, date, time

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

PORT = 8090
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_EXCEL_PATH = os.path.join(BASE_DIR, 'The_Scholars_Database.xlsx')
DB_JSON_PATH = os.path.join(BASE_DIR, 'database.json')

def date_converter(o):
    if isinstance(o, (datetime, date, time)):
        return o.isoformat()
    return str(o)

def read_excel_data():
    if not os.path.exists(DB_EXCEL_PATH):
        if os.path.exists(DB_JSON_PATH):
            with open(DB_JSON_PATH, 'r', encoding='utf-8') as f:
                return json.load(f)
        return {}

    try:
        wb = openpyxl.load_workbook(DB_EXCEL_PATH, data_only=True)
        db = {}
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                db[sheet_name] = []
                continue
            valid_rows = [r for r in rows if any(cell is not None for cell in r)]
            if not valid_rows:
                db[sheet_name] = []
                continue
            headers = [str(h).strip() if h is not None else f'col_{i}' for i, h in enumerate(valid_rows[0])]
            sheet_rows = []
            for r in valid_rows[1:]:
                obj = {}
                for h, val in zip(headers, r):
                    if isinstance(val, (datetime, date, time)):
                        val = val.isoformat()
                    elif isinstance(val, float) and val.is_integer():
                        val = int(val)
                    obj[h] = val
                sheet_rows.append(obj)
            db[sheet_name] = sheet_rows

        wb.close()
        save_json_data(db)
        return db
    except Exception as e:
        print(f"[Warning] Could not read Excel file directly ({e}), loading database.json...")
        if os.path.exists(DB_JSON_PATH):
            with open(DB_JSON_PATH, 'r', encoding='utf-8') as f:
                return json.load(f)
        return {}

def save_json_data(db):
    try:
        with open(DB_JSON_PATH, 'w', encoding='utf-8') as f:
            json.dump(db, f, ensure_ascii=False, indent=2, default=date_converter)
    except Exception as e:
        print(f"[Error] Failed to save database.json: {e}")

def update_excel_user_mode(wa, target_mode):
    if not os.path.exists(DB_EXCEL_PATH):
        return False
    try:
        wb = openpyxl.load_workbook(DB_EXCEL_PATH)
        sheet = wb['STATS']
        wa_norm = str(wa).replace('.0', '').strip()
        found = False

        for row_idx in range(2, sheet.max_row + 1):
            cell_val = sheet.cell(row=row_idx, column=1).value
            if cell_val is not None:
                cell_wa = str(cell_val).replace('.0', '').strip()
                if cell_wa and (wa_norm in cell_wa or cell_wa in wa_norm):
                    sheet.cell(row=row_idx, column=2, value=target_mode)
                    found = True
                    break

        if found:
            wb.save(DB_EXCEL_PATH)
        wb.close()
        return found
    except Exception as e:
        print(f"[Notice] Could not save directly to Excel file ({e}). (Close Excel app if open)")
        return False

def update_excel_global_status(target_status):
    if not os.path.exists(DB_EXCEL_PATH):
        return False
    try:
        wb = openpyxl.load_workbook(DB_EXCEL_PATH)
        sheet = wb['CONFIG']
        found = False

        for row_idx in range(2, sheet.max_row + 1):
            key_val = sheet.cell(row=row_idx, column=1).value
            if key_val is not None and str(key_val).strip() == 'VIRA_STATUS':
                sheet.cell(row=row_idx, column=2, value=target_status)
                found = True
                break

        if not found:
            sheet.append(['VIRA_STATUS', target_status])
            found = True

        wb.save(DB_EXCEL_PATH)
        wb.close()
        return found
    except Exception as e:
        print(f"[Notice] Could not save directly to Excel file ({e}). (Close Excel app if open)")
        return False

class ViraHandler(http.server.BaseHTTPRequestHandler):
    def end_headers_cors(self):
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        self.end_headers()

    def do_OPTIONS(self):
        self.send_response(200)
        self.end_headers_cors()

    def do_GET(self):
        url_path = self.path.split('?')[0]
        if url_path == '/api/data':
            db = read_excel_data()
            self.send_response(200)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.end_headers_cors()
            self.wfile.write(json.dumps({'ok': True, 'database': db}, default=date_converter).encode('utf-8'))
            return

        if url_path == '/' or url_path == '':
            file_path = os.path.join(BASE_DIR, 'index.html')
        else:
            file_path = os.path.join(BASE_DIR, url_path.lstrip('/'))

        if os.path.exists(file_path) and os.path.isfile(file_path):
            mime_type, _ = mimetypes.guess_type(file_path)
            if not mime_type:
                mime_type = 'application/octet-stream'
            self.send_response(200)
            self.send_header('Content-Type', f'{mime_type}; charset=utf-8')
            self.end_headers_cors()
            with open(file_path, 'rb') as f:
                self.wfile.write(f.read())
        else:
            self.send_response(404)
            self.send_header('Content-Type', 'text/plain')
            self.end_headers_cors()
            self.wfile.write(b'File Not Found')

    def do_POST(self):
        content_length = int(self.headers.get('Content-Length', 0))
        post_data = self.rfile.read(content_length)
        
        try:
            payload = json.loads(post_data.decode('utf-8')) if post_data else {}
        except Exception:
            payload = {}

        action = payload.get('action') or self.path.replace('/api/', '')
        db = read_excel_data()

        if action in ['getAllData', 'data']:
            res = {'ok': True, 'database': db}
        elif action in ['setUserMode', 'toggle-user']:
            wa = str(payload.get('wa', '')).strip().replace('.0', '')
            target = str(payload.get('mode', 'ON')).upper()
            found = False

            # 1. Update in-memory db & database.json
            for lead in db.get('STATS', []):
                lead_wa = str(lead.get('No WA') or '').strip().replace('.0', '')
                if lead_wa and (wa in lead_wa or lead_wa in wa):
                    lead['bot_mode'] = target
                    found = True
                    break
            
            save_json_data(db)

            # 2. Update Excel file on disk if possible
            excel_ok = update_excel_user_mode(wa, target)

            res = {
                'ok': True,
                'wa': wa,
                'bot_mode': target,
                'updated': found,
                'excel_updated': excel_ok,
                'message': f"Bot mode for {wa} set to {target}" + (" (Saved to Excel & database.json)" if excel_ok else " (Saved to database.json)")
            }

        elif action in ['setGlobalStatus', 'toggle-global']:
            status = str(payload.get('status', 'ON')).upper()

            # 1. Update in-memory db & database.json
            configs = db.get('CONFIG', [])
            cfg = next((c for c in configs if c.get('Key') == 'VIRA_STATUS'), None)
            if cfg:
                cfg['Value'] = status
            else:
                configs.append({'Key': 'VIRA_STATUS', 'Value': status})
            db['CONFIG'] = configs

            save_json_data(db)

            # 2. Update Excel file on disk if possible
            excel_ok = update_excel_global_status(status)

            res = {
                'ok': True,
                'status': status,
                'excel_updated': excel_ok,
                'message': f"Global VIRA status set to {status}" + (" (Saved to Excel & database.json)" if excel_ok else " (Saved to database.json)")
            }
        else:
            res = {'ok': False, 'error': 'UNKNOWN_ACTION', 'message': f'Unknown action: {action}'}

        self.send_response(200)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.end_headers_cors()
        self.wfile.write(json.dumps(res, default=date_converter).encode('utf-8'))

class ReusableServer(socketserver.TCPServer):
    allow_reuse_address = True

if __name__ == '__main__':
    print("==================================================")
    print("VIRA Performance & CRM Dashboard Server Started!")
    print(f"URL: http://localhost:{PORT}")
    print("==================================================")
    with ReusableServer(("", PORT), ViraHandler) as httpd:
        httpd.serve_forever()
