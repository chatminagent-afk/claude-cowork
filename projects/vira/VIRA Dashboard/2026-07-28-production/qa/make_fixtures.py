#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
make_fixtures.py — membuat data uji untuk selftest.html dan mock-api.py.

Prinsipnya: **nama kolom diambil dari spreadsheet produksi yang sebenarnya**,
tapi seluruh isinya sintetis. Jadi kalau ada kolom yang berubah/ditulis dengan
trailing space di produksi, test ikut memakai bentuk yang sama — sementara
tidak ada satu pun nomor WhatsApp atau nama asli yang keluar dari mesin ini.

Jalankan ulang kalau struktur spreadsheet berubah:
    python qa/make_fixtures.py
"""

import json
import os
import random
import sys
from datetime import datetime, timedelta

try:
    import openpyxl
except ImportError:
    print('butuh openpyxl:  pip install openpyxl')
    sys.exit(1)

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'fixtures')
ROOT = os.path.abspath(os.path.join(HERE, '..', '..', '..', '..'))

BOOKS = {
    'thescholars': os.path.join(ROOT, 'the scholars', 'report', 'production',
                                'The_Scholars_Database.xlsx'),
    'persada': os.path.join(ROOT, 'Persada Cisoka Residence', 'workflow', 'production',
                            'PCR_Database.xlsx'),
}

TABS = {
    'thescholars': ['STATS', 'UNKNOWN', 'MOCK_INTERVIEW_BOOKING', 'MOCK_SLOTS'],
    'persada': ['STATS', 'UNKNOWN', 'SURVEY', 'EVENTS'],
}

# Titik acuan waktu test. Dibekukan supaya hasil test deterministik —
# selftest.html memakai nilai yang sama sebagai `nowMs`.
NOW = datetime(2026, 7, 28, 10, 0, 0)
NOW_MS = int(NOW.timestamp() * 1000)

random.seed(20260728)

NAMA = ['Andi', 'Budi', 'Citra', 'Dewi', 'Eka', 'Fajar', 'Gita', 'Hadi', 'Indah',
        'Joko', 'Kirana', 'Lukman', 'Maya', 'Nanda', 'Oki', 'Putri', 'Rian',
        'Sari', 'Tono', 'Umi', 'Vina', 'Wawan', 'Yuni', 'Zaki']

# Rentang nomor DIBEDAKAN per tenant. Kalau keduanya memakai rentang yang sama,
# uji "toggle nomor milik tenant lain harus ditolak" jadi tidak bermakna —
# nomornya kebetulan ada di kedua sisi dan toggle-nya wajar saja berhasil.
PREFIX = {'thescholars': '62812', 'persada': '62813'}
LID_EDGE = {'thescholars': ('199988877766655', '288776655443322', '62899000111'),
            'persada': ('399988877766655', '488776655443322', '62899000222')}


def headers_of(path, tab):
    """Baca baris header persis apa adanya — termasuk trailing space kalau ada."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if tab not in wb.sheetnames:
        wb.close()
        return None
    ws = wb[tab]
    cols = []
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        for c in row:
            if c is None:
                continue
            cols.append(str(c))
        break
    wb.close()
    return cols


def dstr(dt):
    return dt.strftime('%d/%m/%Y')


def blank_row(cols):
    return dict((c, '') for c in cols)


def gen_stats(cols, tenant, n):
    """
    Baris STATS sintetis + sejumlah kasus tepi yang HARUS ditangani kode:
      - identitas hanya LID (No WA kosong)
      - nomor 15+ digit (LID yang tersimpan di kolom No WA)
      - bot_mode huruf kecil / kosong
      - tanggal tidak valid
      - Counter berformat ribuan
      - baris kosong total
    """
    rows = []
    for i in range(n):
        c = blank_row(cols)
        first = NOW - timedelta(days=random.randint(0, 100))
        last = first + timedelta(days=random.randint(0, 20))
        if last > NOW:
            last = NOW
        wa = PREFIX[tenant] + '%07d' % (1000000 + i)
        c['No WA'] = wa
        c['lid'] = '%015d' % (100000000000000 + i)
        c['Nama'] = random.choice(NAMA) + ' ' + chr(65 + (i % 26)) + '.'
        c['bot_mode'] = random.choice(['ON', 'ON', 'ON', 'OFF', ''])
        c['Counter'] = str(random.randint(1, 60))
        c['Intensitas Chat'] = random.choice(['Rendah', 'Sedang', 'Tinggi'])
        c['Tanggal Chat Pertama'] = dstr(first)
        c['Tanggal Chat Terakhir'] = dstr(last)
        c['Jam Chat Terakhir'] = '%02d:%02d:00' % (random.randint(6, 22), random.randint(0, 59))
        c['Pesan Pertama'] = 'halo kak mau tanya'
        c['greeting_sent'] = random.choice(['Y', ''])
        c['follow_up_count'] = str(random.randint(0, 3))

        if tenant == 'thescholars':
            c['program_interest'] = random.choice(['Junior', 'Intermediate', 'Seniors', ''])
            c['kelas_anak'] = random.choice(['SMP 1', 'SMP 2', 'SMA 10', 'SMA 11', 'SMA 12',
                                             'SMA 11, SMA 12', ''])
            c['gform_sent_ts'] = random.choice([str(int(last.timestamp())), ''])
            c['gform_filled'] = ''      # memang tidak pernah terisi di produksi
        else:
            c['nama_lengkap'] = c['Nama'] + 'saputra'
            c['lead_source'] = random.choice(['Instagram', 'Facebook', 'Google', 'TikTok',
                                              'WhatsApp', 'Organik', ''])
            c['unit_interest'] = random.choice(['Tipe 36/72', 'Tipe 36/81', 'Tipe 30/60', ''])
            c['budget_range'] = random.choice(['300-400jt', '400-500jt', ''])
            c['lokasi_kerja'] = random.choice(['Tangerang', 'Jakarta Barat', ''])
            if random.random() < 0.3:
                c['survey_date'] = dstr(last + timedelta(days=random.randint(1, 7)))
                c['survey_time'] = '10:00'
                c['survey_status'] = random.choice(['SCHEDULED', 'DONE', 'CANCELLED'])
                c['flag_survey'] = 'Y'
            else:
                c['survey_status'] = ''
        rows.append(c)

    # --- kasus tepi eksplisit ---
    e = blank_row(cols)                       # hanya LID, No WA kosong
    e['No WA'] = ''
    e['lid'] = LID_EDGE[tenant][0]
    e['Nama'] = 'User Tanpa Nomor'
    e['bot_mode'] = 'off'                     # huruf kecil -> tetap harus dibaca OFF
    e['Counter'] = '7'
    e['Tanggal Chat Pertama'] = dstr(NOW - timedelta(days=3))
    e['Tanggal Chat Terakhir'] = dstr(NOW - timedelta(days=1))
    e['Jam Chat Terakhir'] = '25:00:00'       # jam mustahil -> harus diabaikan
    rows.append(e)

    e2 = blank_row(cols)                      # LID tersimpan di kolom No WA
    e2['No WA'] = LID_EDGE[tenant][1]
    e2['Nama'] = 'User LID Lama'
    e2['bot_mode'] = ''
    e2['Counter'] = '1.234'                   # format ribuan
    e2['Tanggal Chat Pertama'] = '31/13/2026'  # tanggal tidak valid
    e2['Tanggal Chat Terakhir'] = dstr(NOW)
    e2['Jam Chat Terakhir'] = '09:15:00'
    rows.append(e2)

    rows.append(blank_row(cols))              # baris kosong total -> harus dibuang

    e3 = blank_row(cols)                      # lead lama, di luar semua rentang
    e3['No WA'] = LID_EDGE[tenant][2]
    e3['Nama'] = 'Lead Lama'
    e3['bot_mode'] = 'ON'
    e3['Counter'] = '2'
    e3['Tanggal Chat Pertama'] = dstr(NOW - timedelta(days=200))
    e3['Tanggal Chat Terakhir'] = dstr(NOW - timedelta(days=180))
    e3['Jam Chat Terakhir'] = '13:00:00'
    rows.append(e3)

    for idx, r in enumerate(rows):
        r['row_number'] = idx + 2             # seperti output node Google Sheets
    return rows


def gen_unknown(cols, n, prefix):
    rows = []
    for i in range(n):
        c = blank_row(cols)
        d = NOW - timedelta(days=random.randint(0, 120))
        c['Tanggal'] = dstr(d)
        c['User'] = prefix + '%07d' % (1000000 + random.randint(0, 40))
        c['Pertanyaan'] = 'pertanyaan uji nomor %d' % (i + 1)
        c['message'] = ''
        c['row_number'] = i + 2
        rows.append(c)
    return rows


def gen_simple(cols, n, filler):
    rows = []
    for i in range(n):
        c = blank_row(cols)
        filler(c, i)
        c['row_number'] = i + 2
        rows.append(c)
    return rows


def build(tenant):
    path = BOOKS[tenant]
    if not os.path.exists(path):
        print('[!] tidak ketemu: %s' % path)
        return None
    data = {}
    for tab in TABS[tenant]:
        cols = headers_of(path, tab)
        if cols is None:
            print('[!] tab %s tidak ada di %s' % (tab, tenant))
            data[tab] = []
            continue

        if tab == 'STATS':
            data[tab] = gen_stats(cols, tenant, 60 if tenant == 'thescholars' else 40)
        elif tab == 'UNKNOWN':
            data[tab] = gen_unknown(cols, 25 if tenant == 'thescholars' else 12,
                                    PREFIX[tenant])
        elif tab == 'MOCK_INTERVIEW_BOOKING':
            def f(c, i):
                c['ID Booking'] = 'MB-%03d' % (i + 1)
                c['Tanggal'] = dstr(NOW - timedelta(days=i * 3))
                c['Sesi'] = random.choice(['Sesi 1', 'Sesi 2'])
                c['Nama Anak'] = random.choice(NAMA)
                c['Nama Orang Tua'] = random.choice(NAMA)
                c['Sekolah'] = 'SMA Negeri %d' % (i + 1)
                c['Kelas'] = random.choice(['10', '11', '12'])
                c['Target Uni Jurusan'] = random.choice(['NUS CS', 'NTU Business', 'SMU Law'])
                c['Status'] = random.choice(['Confirmed', 'Pending', 'Done'])
            data[tab] = gen_simple(cols, 9, f)
        elif tab == 'MOCK_SLOTS':
            def f(c, i):
                c['ID'] = 'SL-%03d' % (i + 1)
                c['Tanggal'] = dstr(NOW + timedelta(days=i))
                c['JamMulai'] = '%02d:00' % (9 + (i % 6))
                c['JamSelesai'] = '%02d:00' % (10 + (i % 6))
                c['Booked'] = random.choice(['Y', ''])
                c['NamaPemesan'] = random.choice(NAMA) if random.random() < 0.5 else ''
            data[tab] = gen_simple(cols, 14, f)
        elif tab == 'SURVEY':
            def f(c, i):
                d = NOW - timedelta(days=i)
                c['no_wa'] = PREFIX[tenant] + '%07d' % (1000000 + i)
                c['nama'] = random.choice(NAMA)
                c['tanggal'] = dstr(d + timedelta(days=2))
                c['jam'] = '%02d:00' % (9 + (i % 7))
                c['unit_diminati'] = random.choice(['Tipe 36/72', 'Tipe 36/81'])
                c['status'] = random.choice(['SCHEDULED', 'DONE'])
                c['sumber_traffic'] = random.choice(['Instagram', 'Google', 'Facebook'])
                c['catatan'] = 'ringkasan uji %d' % (i + 1)
                c['created_ts'] = str(int(d.timestamp()))
            data[tab] = gen_simple(cols, 11, f)
        elif tab == 'EVENTS':
            def f(c, i):
                d = NOW - timedelta(days=i)
                c['no_wa'] = PREFIX[tenant] + '%07d' % (1000000 + i)
                c['nama'] = random.choice(NAMA)
                c['event'] = random.choice(['DELEGATED', 'REQUEST_CALL'])
                c['detail'] = 'detail uji %d' % (i + 1)
                # SENGAJA campur satuan: detik untuk DELEGATED, milidetik untuk
                # REQUEST_CALL — persis inkonsistensi yang ada di produksi.
                c['ts'] = str(int(d.timestamp())) if c['event'] == 'DELEGATED' \
                    else str(int(d.timestamp() * 1000))
            data[tab] = gen_simple(cols, 8, f)
        else:
            data[tab] = []
    return data


def main():
    if not os.path.isdir(OUT):
        os.makedirs(OUT)
    manifest = {'now_ms': NOW_MS, 'generated_for': '2026-07-28', 'tenants': {}}
    for tenant in BOOKS:
        data = build(tenant)
        if data is None:
            return 1
        p = os.path.join(OUT, '%s.json' % tenant)
        with open(p, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=1)
        counts = dict((k, len(v)) for k, v in data.items())
        manifest['tenants'][tenant] = counts
        print('[ok] fixtures/%s.json  %s' % (tenant, counts))

    with open(os.path.join(OUT, 'manifest.json'), 'w', encoding='utf-8') as f:
        json.dump(manifest, f, ensure_ascii=False, indent=1)
    print('[ok] fixtures/manifest.json  now_ms=%d' % NOW_MS)
    return 0


if __name__ == '__main__':
    sys.exit(main())
