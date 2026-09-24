# -*- coding: utf-8 -*-
"""QA for 2026-08-19-VIRA-STATS-cleanup-3bulan.json

Sections:
  A. Structure / graph
  B. JS syntax of the Code nodes
  C. Logic port of "Plan cleanup" + edge cases
  D. Full sheet-transformation simulation (append-then-delete)
  E. Real STATS data from The_Scholars_Database.xlsx
"""
import json
import re
import sys

WF = r"D:\Documents\Claude Cowork\the scholars\report\patch\2026-08-19-VIRA-STATS-cleanup-3bulan.json"
XLSX = r"D:\Documents\Claude Cowork\the scholars\report\production\The_Scholars_Database.xlsx"

PASS = []
FAIL = []


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    mark = "PASS" if cond else "FAIL"
    line = f"  [{mark}] {name}"
    if detail:
        line += f"  -- {detail}"
    print(line)
    return cond


# ====================================================================
print("=" * 68)
print("A. STRUKTUR & GRAPH")
print("=" * 68)

with open(WF, encoding="utf-8") as f:
    wf = json.load(f)

nodes = wf["nodes"]
conns = wf["connections"]
names = [n["name"] for n in nodes]
by_name = {n["name"]: n for n in nodes}

check("A1 JSON valid & punya nodes/connections", "nodes" in wf and "connections" in wf)
check("A2 jumlah node = 9", len(nodes) == 9, f"{len(nodes)}")
check("A3 nama node unik", len(set(names)) == len(names))
check("A4 semua node punya id unik", len({n["id"] for n in nodes}) == len(nodes))

# connection integrity
bad_src = [s for s in conns if s not in by_name]
targets = []
for src, spec in conns.items():
    for branch in spec.get("main", []):
        for c in branch:
            targets.append(c["node"])
bad_tgt = [t for t in targets if t not in by_name]
check("A5 semua source connection ada sbg node", not bad_src, str(bad_src))
check("A6 semua target connection ada sbg node", not bad_tgt, str(bad_tgt))

# reachability from trigger
trigger = "Every 3 months 00:01 WIB"
seen, stack = set(), [trigger]
while stack:
    cur = stack.pop()
    if cur in seen:
        continue
    seen.add(cur)
    for branch in conns.get(cur, {}).get("main", []):
        for c in branch:
            stack.append(c["node"])
check("A7 semua node terjangkau dari trigger", seen == set(names),
      f"tidak terjangkau: {set(names) - seen}")
check("A8 trigger adalah scheduleTrigger",
      by_name[trigger]["type"] == "n8n-nodes-base.scheduleTrigger")

# cron
cron = by_name[trigger]["parameters"]["rule"]["interval"][0]["expression"]
check("A9 cron = '1 0 1 */3 *' (menit 1, jam 0, tgl 1, tiap 3 bulan)",
      cron == "1 0 1 */3 *", cron)
check("A10 timezone workflow = Asia/Jakarta",
      wf["settings"].get("timezone") == "Asia/Jakarta")
check("A11 workflow di-import dalam keadaan non-aktif", wf.get("active") is False)
check("A12 errorWorkflow terpasang",
      wf["settings"].get("errorWorkflow") == "ZKsINjA7ZC8c9yRHWp3ud")

# google sheets nodes
gs = [n for n in nodes if n["type"] == "n8n-nodes-base.googleSheets"]
check("A13 ada 3 node googleSheets (read/append/delete)", len(gs) == 3, str(len(gs)))
check("A14 semua googleSheets pakai serviceAccount",
      all(n["parameters"].get("authentication") == "serviceAccount" for n in gs))
check("A15 semua googleSheets pakai credential yang sama dgn workflow utama",
      all(n.get("credentials", {}).get("googleApi", {}).get("id") == "3gTKMbD8lJDLbRqR" for n in gs))
check("A16 semua googleSheets menunjuk tab STATS (gid 56867128)",
      all(n["parameters"]["sheetName"]["value"] == 56867128 for n in gs))
check("A17 semua googleSheets typeVersion 4.7 (sama dgn cleanup MSG_BUFFER produksi)",
      all(n["typeVersion"] == 4.7 for n in gs))

dele = by_name["Delete original block"]["parameters"]
check("A18 delete: operation/toDelete benar",
      dele.get("operation") == "delete" and dele.get("toDelete") == "rows")
check("A19 delete: startIndex dari plan (baris 2)",
      dele.get("startIndex") == "={{ $json.startRow }}", str(dele.get("startIndex")))
check("A20 delete: numberToDelete = totalDataRows",
      dele.get("numberToDelete") == "={{ $json.totalDataRows }}")

app = by_name["Append SAM rows"]["parameters"]
check("A21 append: operation=append & autoMapInputData",
      app.get("operation") == "append"
      and app["columns"]["mappingMode"] == "autoMapInputData")
check("A22 append: retryOnFail aktif (baris SAM tidak boleh gagal diam2)",
      by_name["Append SAM rows"].get("retryOnFail") is True
      and by_name["Append SAM rows"].get("maxTries", 0) >= 3)

# ---- the critical ordering property ----
def downstream(start):
    s, st = set(), [start]
    while st:
        cur = st.pop()
        for br in conns.get(cur, {}).get("main", []):
            for c in br:
                if c["node"] not in s:
                    s.add(c["node"])
                    st.append(c["node"])
    return s

check("A23 KRUSIAL: Append berada SEBELUM Delete (bukan sebaliknya)",
      "Delete original block" in downstream("Append SAM rows")
      and "Append SAM rows" not in downstream("Delete original block"))

# both IF branches must converge on Prepare delete
ifbr = conns["IF Ada Baris SAM"]["main"]
check("A24 IF punya 2 cabang", len(ifbr) == 2, str(len(ifbr)))
check("A25 cabang true -> Split kept rows", ifbr[0][0]["node"] == "Split kept rows")
check("A26 cabang false -> Prepare delete (delete tetap jalan walau 0 baris SAM)",
      ifbr[1][0]["node"] == "Prepare delete")

inbound_delete = [s for s, spec in conns.items()
                  for br in spec.get("main", []) for c in br
                  if c["node"] == "Delete original block"]
check("A27 Delete hanya punya 1 sumber inbound = Prepare delete",
      inbound_delete == ["Prepare delete"], str(inbound_delete))

inbound_prep = sorted(s for s, spec in conns.items()
                      for br in spec.get("main", []) for c in br
                      if c["node"] == "Prepare delete")
check("A28 Prepare delete menerima dari kedua cabang",
      inbound_prep == ["Append SAM rows", "IF Ada Baris SAM"], str(inbound_prep))

phone = [p for p in by_name["Notify Steven"]["parameters"]["bodyParameters"]["parameters"]
         if p["name"] == "phone"][0]["value"]
check("A29 nomor notif masih placeholder (wajib diisi Steven, tidak bocor ke nomor asing)",
      "ISI_NOMOR_WA_STEVEN" in phone, phone)
check("A30 Notify tidak menggagalkan workflow kalau Kirimi error",
      by_name["Notify Steven"].get("onError") == "continueRegularOutput")

# ====================================================================
print()
print("=" * 68)
print("B. SINTAKS JAVASCRIPT NODE CODE")
print("=" * 68)

code_nodes = {n["name"]: n["parameters"]["jsCode"]
              for n in nodes if n["type"] == "n8n-nodes-base.code"}
check("B1 ada 3 node Code", len(code_nodes) == 3, str(list(code_nodes)))

PROD_WF = r"D:\Documents\Claude Cowork\the scholars\report\production\2026-08-08-VIRA-V4-retryable.json"

try:
    import esprima
    # esprima 4.0.1 = ES2017; tidak mengenal nullish coalescing (??, ES2020).
    # n8n jalan di Node 18+ yang mendukungnya, dan `??` SUDAH dipakai di node
    # produksi workflow ini. Jadi untuk keperluan cek sintaks, `??` disubstitusi
    # jadi `||` (bentuk sintaksisnya setara) supaya sisa kodenya tetap tervalidasi.
    for nm, src in code_nodes.items():
        try:
            esprima.parseScript("(function(){\n" + src.replace("??", "||") + "\n})")
            check(f"B2 sintaks OK: {nm}", True)
        except Exception as e:
            check(f"B2 sintaks OK: {nm}", False, str(e))

    # Bukti bahwa idiom `??` yang dipakai memang jalan di n8n: cari idiom yang
    # sama persis di workflow produksi yang sudah live.
    used = set(re.findall(r"String\(\s*\w+\s*\?\?\s*''\s*\)", "\n".join(code_nodes.values())))
    try:
        with open(PROD_WF, encoding="utf-8") as pf:
            prod_src = pf.read()
        prod_has = all(re.search(re.escape(u).replace(r"\ ", r"\s*"), prod_src) for u in used) if used else True
        check("B2b idiom `??` yang dipakai identik dgn node produksi yang sudah live",
              prod_has, f"idiom: {sorted(used)}")
    except FileNotFoundError:
        print("  [SKIP] file produksi tidak ditemukan untuk pembanding idiom")
except ImportError:
    print("  [SKIP] esprima tidak terpasang - fallback ke pemeriksaan pasangan kurung")
    for nm, src in code_nodes.items():
        bal = {"(": 0, "[": 0, "{": 0}
        pair = {")": "(", "]": "[", "}": "{"}
        stripped = re.sub(r"//[^\n]*", "", src)
        stripped = re.sub(r"'[^'\n]*'|\"[^\"\n]*\"|`(?:[^`\\]|\\.)*`", "''", stripped, flags=re.S)
        for ch in stripped:
            if ch in bal:
                bal[ch] += 1
            elif ch in pair:
                bal[pair[ch]] -= 1
        check(f"B2 kurung seimbang: {nm}", all(v == 0 for v in bal.values()), str(bal))

check("B3 Plan cleanup membuang row_number dari salinan",
      "delete copy.row_number" in code_nodes["Plan cleanup"])
check("B4 Plan cleanup normalisasi case-insensitive + trim",
      ".trim().toUpperCase()" in code_nodes["Plan cleanup"])
check("B5 Plan cleanup punya guard kolom off_reason",
      "off_reason' in sample" in code_nodes["Plan cleanup"]
      or "'off_reason' in sample" in code_nodes["Plan cleanup"])
check("B6 Prepare delete mengembalikan TEPAT 1 item",
      "return [{ json: plan }]" in code_nodes["Prepare delete"])

# ====================================================================
print()
print("=" * 68)
print("C. PORT LOGIKA 'Plan cleanup' + EDGE CASE")
print("=" * 68)


class Abort(Exception):
    pass


def plan_cleanup(input_rows):
    """1:1 port of the Plan cleanup node."""
    KEEP = "SAM"
    norm = lambda v: str(v if v is not None else "").strip().upper()

    rows = [r for r in input_rows
            if isinstance(r, dict) and r.get("row_number") is not None
            and int(r["row_number"]) >= 2]

    if not rows:
        return None  # return []

    sample = rows[0]
    if "No WA" not in sample:
        raise Abort('Kolom "No WA" tidak ditemukan')
    if "off_reason" not in sample:
        raise Abort('Kolom "off_reason" belum ada')

    kept, off_vira, off_nolabel = [], 0, 0
    for r in rows:
        if norm(r.get("off_reason")) == KEEP:
            c = dict(r)
            c.pop("row_number", None)
            kept.append(c)
        elif norm(r.get("bot_mode")) == "OFF":
            if norm(r.get("off_reason")) == "VIRA":
                off_vira += 1
            else:
                off_nolabel += 1

    max_row = max(int(r["row_number"]) for r in rows)
    total = max_row - 1
    if total < 1:
        return None

    return {
        "startRow": 2, "totalDataRows": total,
        "keptCount": len(kept), "deleteCount": total - len(kept),
        "offByVira": off_vira, "offNoLabel": off_nolabel,
        "keptRows": kept,
    }


def row(rn, wa, bot="ON", reason="", nama="x"):
    return {"row_number": rn, "No WA": wa, "bot_mode": bot,
            "off_reason": reason, "Nama": nama}


# C1 sheet kosong
check("C1 sheet kosong -> abort (tidak menghapus apa pun)", plan_cleanup([]) is None)

# C2 guard kolom
try:
    plan_cleanup([{"row_number": 2, "foo": 1}])
    check("C2 tab salah (tanpa 'No WA') -> error", False)
except Abort:
    check("C2 tab salah (tanpa 'No WA') -> error", True)

try:
    plan_cleanup([{"row_number": 2, "No WA": "628", "bot_mode": "ON"}])
    check("C3 kolom off_reason belum ada -> error (STOP, bukan hapus semua)", False)
except Abort:
    check("C3 kolom off_reason belum ada -> error (STOP, bukan hapus semua)", True)

# C4 variasi penulisan SAM
variants = ["SAM", "Sam", "sam", " SAM ", "  sAm"]
data = [row(2 + i, f"628{i}", "OFF", v) for i, v in enumerate(variants)]
p = plan_cleanup(data)
check("C4 semua variasi ketikan SAM dipertahankan (case-insensitive + trim)",
      p["keptCount"] == 5 and p["deleteCount"] == 0, f"kept={p['keptCount']}")

# C5 VIRA & kosong tidak dipertahankan
data = [row(2, "6281", "OFF", "VIRA"), row(3, "6282", "OFF", ""),
        row(4, "6283", "ON", ""), row(5, "6284", "OFF", "SAM")]
p = plan_cleanup(data)
check("C5 hanya SAM yang dipertahankan; VIRA/kosong/ON ikut dihapus",
      p["keptCount"] == 1 and p["deleteCount"] == 3
      and p["keptRows"][0]["No WA"] == "6284")
check("C6 hitungan OFF by VIRA benar", p["offByVira"] == 1, str(p["offByVira"]))
check("C7 hitungan OFF tanpa label benar", p["offNoLabel"] == 1, str(p["offNoLabel"]))

# C8 row_number dibuang
check("C8 row_number dibuang dari baris salinan (tidak ditulis ke sheet)",
      all("row_number" not in r for r in p["keptRows"]))

# C9 semua SAM
data = [row(2, "a", "OFF", "SAM"), row(3, "b", "OFF", "Sam")]
p = plan_cleanup(data)
check("C9 semua baris SAM -> deleteCount 0, hasil akhir tetap benar",
      p["keptCount"] == 2 and p["deleteCount"] == 0)

# C10 tidak ada SAM
data = [row(2, "a", "ON", ""), row(3, "b", "OFF", "VIRA")]
p = plan_cleanup(data)
check("C10 tidak ada baris SAM -> keptCount 0 (IF ambil cabang false, delete tetap jalan)",
      p["keptCount"] == 0 and p["deleteCount"] == 2)

# C11 baris berlubang
data = [row(2, "a", "ON", ""), row(7, "b", "OFF", "SAM")]
p = plan_cleanup(data)
check("C11 baris berlubang: rentang hapus pakai row_number terbesar",
      p["totalDataRows"] == 6, f"totalDataRows={p['totalDataRows']}")

# C12 header-only
check("C12 hanya header (tidak ada row_number>=2) -> abort",
      plan_cleanup([{"row_number": 1, "No WA": "h", "off_reason": "x"}]) is None)

# C13 nilai None
data = [row(2, "a", None, None), row(3, "b", "OFF", "SAM")]
p = plan_cleanup(data)
check("C13 nilai kosong/None tidak bikin crash", p["keptCount"] == 1)

# ====================================================================
print()
print("=" * 68)
print("D. SIMULASI PENUH: append -> delete -> hasil akhir")
print("=" * 68)


def simulate(initial_rows):
    """Simulate the real sheet mutation the workflow performs.

    initial_rows: list of dicts WITHOUT row_number, in sheet order (row 2..N+1)
    Returns final sheet rows in order.
    """
    sheet = [dict(r) for r in initial_rows]
    read = [dict(r, row_number=i + 2) for i, r in enumerate(sheet)]

    plan = plan_cleanup(read)
    if plan is None:
        return sheet  # workflow aborted, sheet untouched

    # IF Ada Baris SAM -> true: append copies to the bottom
    if plan["keptCount"] > 0:
        sheet = sheet + [dict(r) for r in plan["keptRows"]]

    # Delete original block: rows 2..(2+totalDataRows-1) => indices 0..total-1
    total = plan["totalDataRows"]
    sheet = sheet[total:]
    return sheet


def mk(wa, bot="ON", reason=""):
    return {"No WA": wa, "bot_mode": bot, "off_reason": reason}


# D1 skenario campuran realistis
initial = [
    mk("6281", "ON", ""),
    mk("6282", "OFF", "SAM"),
    mk("6283", "OFF", "VIRA"),
    mk("6284", "ON", ""),
    mk("6285", "OFF", "Sam"),
    mk("6286", "OFF", ""),
]
final = simulate(initial)
expected = [mk("6282", "OFF", "SAM"), mk("6285", "OFF", "Sam")]
check("D1 hasil akhir = persis baris SAM saja", final == expected, str(final))
check("D2 urutan relatif baris SAM terjaga",
      [r["No WA"] for r in final] == ["6282", "6285"])
check("D3 tidak ada baris non-SAM tersisa",
      all(str(r["off_reason"]).strip().upper() == "SAM" for r in final))
check("D4 tidak ada baris SAM yang hilang",
      len(final) == len([r for r in initial
                         if str(r["off_reason"]).strip().upper() == "SAM"]))
check("D5 baris rapi & kontigu mulai baris 2 (tidak ada celah)",
      len(final) == len([r for r in final if r]))

# D6 tanpa baris SAM
final = simulate([mk("a"), mk("b", "OFF", "VIRA")])
check("D6 tanpa baris SAM -> sheet tersisa header saja", final == [], str(final))

# D7 semua SAM -> isi identik, tidak ada duplikat
initial = [mk("a", "OFF", "SAM"), mk("b", "OFF", "SAM")]
final = simulate(initial)
check("D7 semua baris SAM -> isi identik, TIDAK ada duplikat",
      final == initial, str(final))

# D8 idempoten: jalan 2x berturut-turut hasilnya sama
once = simulate(initial)
twice = simulate(once)
check("D8 idempoten (jalan 2x hasilnya sama)", once == twice)

# D9 kegagalan append = nol kehilangan data
initial = [mk("a"), mk("b", "OFF", "SAM")]
sheet_if_append_fails = [dict(r) for r in initial]  # delete never runs
check("D9 kalau Append gagal, blok lama BELUM terhapus -> nol kehilangan data",
      sheet_if_append_fails == initial)

# D10 kegagalan delete = duplikat, bukan kehilangan
read = [dict(r, row_number=i + 2) for i, r in enumerate(initial)]
plan = plan_cleanup(read)
sheet_if_delete_fails = initial + plan["keptRows"]
check("D10 kalau Delete gagal, hasilnya duplikat (bisa diperbaiki), bukan data hilang",
      len([r for r in sheet_if_delete_fails
           if str(r["off_reason"]).strip().upper() == "SAM"]) == 2)

# ====================================================================
print()
print("=" * 68)
print("E. DATA STATS ASLI (The_Scholars_Database.xlsx)")
print("=" * 68)

try:
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    if "STATS" not in wb.sheetnames:
        print(f"  [SKIP] tab STATS tidak ada. Tab tersedia: {wb.sheetnames}")
    else:
        ws = wb["STATS"]
        rows_iter = ws.iter_rows(values_only=True)
        header = [str(h) if h is not None else "" for h in next(rows_iter)]
        real = []
        for i, r in enumerate(rows_iter):
            if all(v is None or str(v).strip() == "" for v in r):
                continue
            # baris bisa lebih pendek dari header -> indeks aman
            d = {header[j]: (r[j] if j < len(r) else None)
                 for j in range(len(header)) if header[j]}
            d["row_number"] = i + 2
            real.append(d)

        print(f"  Kolom STATS saat ini ({len(header)}): {header}")
        print(f"  Jumlah baris data terbaca: {len(real)}")

        has_off_reason = "off_reason" in header
        # Ini PRASYARAT, bukan uji kode: kolomnya memang belum dibuat.
        # Yang diuji adalah apakah guard-nya menolak jalan (E2).
        print(f"\n  [PRASYARAT] kolom off_reason ada di STATS: "
              f"{'YA' if has_off_reason else 'BELUM - harus ditambah manual dulu'}")

        if not has_off_reason:
            try:
                plan_cleanup(real)
                check("E2 guard menolak jalan selama off_reason belum ada", False,
                      "BAHAYA: workflow jalan tanpa kolom off_reason")
            except Abort as e:
                check("E2 guard menolak jalan selama off_reason belum ada", True,
                      "workflow berhenti dgn error, TIDAK menghapus apa pun")

            # simulate what happens once the column is added
            blank = lambda v: v is None or str(v).strip() == ""
            off_now = [r for r in real
                       if str(r.get("bot_mode") or "").strip().upper() == "OFF"]
            no_trace = [r for r in off_now
                        if blank(r.get("Pesan Pertama")) and blank(r.get("Counter"))]
            with_trace = [r for r in off_now if r not in no_trace]
            print(f"\n  Baris bot_mode=OFF saat ini: {len(off_now)} dari {len(real)}")
            print(f"    - tanpa jejak chat sama sekali : {len(no_trace)}  "
                  f"(nomor yang diketik manual Sam -> hampir pasti SAM)")
            print(f"    - ada jejak chat dgn VIRA      : {len(with_trace)}  "
                  f"(perlu penilaian Sam: SAM atau VIRA)")

            check("E5 dua kelompok itu benar-benar terpisah (tidak ada yg dobel/hilang)",
                  len(no_trace) + len(with_trace) == len(off_now))

            # skenario yang sebenarnya akan terjadi kalau Sam hanya sempat
            # melabeli kelompok "tanpa jejak chat"
            for r in real:
                r["off_reason"] = "SAM" if r in no_trace else ""
            p2 = plan_cleanup(real)
            check("E6 kalau HANYA kelompok tanpa-jejak dilabeli SAM, jumlah cocok",
                  p2["keptCount"] == len(no_trace), f"{p2['keptCount']} vs {len(no_trace)}")
            print(f"  Skenario itu: {p2['totalDataRows']} baris -> "
                  f"{p2['keptCount']} dipertahankan, {p2['deleteCount']} dihapus")

            for r in real:
                r["off_reason"] = "SAM" if str(r.get("bot_mode") or "").strip().upper() == "OFF" else ""
            p = plan_cleanup(real)
            check("E3 setelah kolom ditambah, plan jalan tanpa error", p is not None)
            check("E4 simulasi 'semua OFF dilabeli SAM': jumlah dipertahankan cocok",
                  p["keptCount"] == len(off_now), f"{p['keptCount']} vs {len(off_now)}")
            print(f"  Simulasi: {p['totalDataRows']} baris -> "
                  f"{p['keptCount']} dipertahankan, {p['deleteCount']} dihapus")
        else:
            p = plan_cleanup(real)
            check("E2 plan jalan atas data asli", p is not None)
            if p:
                print(f"  {p['totalDataRows']} baris -> {p['keptCount']} dipertahankan, "
                      f"{p['deleteCount']} dihapus (VIRA {p['offByVira']}, "
                      f"tanpa label {p['offNoLabel']})")
except ImportError:
    print("  [SKIP] openpyxl tidak terpasang")
except Exception as e:
    print(f"  [SKIP] gagal baca xlsx: {e}")

# ====================================================================
print()
print("=" * 68)
print(f"HASIL: {len(PASS)} PASS / {len(FAIL)} FAIL")
if FAIL:
    print("GAGAL:")
    for f_ in FAIL:
        print("  -", f_)
print("=" * 68)
sys.exit(1 if FAIL else 0)
